from pathlib import Path
import fitz
from pptx import Presentation
from backend.app.config import Settings
from backend.app.ingestion import extract, ingest, chunk_record
from backend.app.store import Store


def test_pdf_metadata_and_duplicate_ingestion(tmp_path):
    pdf = fitz.open()
    for text in ['Feedback compares the measured output with the desired reference input.', 'Negative feedback reduces the error signal in the closed loop system.']:
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(tmp_path / 'lecture.pdf')
    db = Store(tmp_path / 'test.db')
    result = ingest(tmp_path, db, Settings(embedding_provider='none'))
    assert len(result['ingested']) == 1
    assert {c.page for c in db.chunks()} == {1, 2}
    assert len(ingest(tmp_path, db, Settings())['skipped']) == 1
    assert len(db.chunks()) == 2


def test_slides_and_markdown(tmp_path):
    deck = Presentation()
    for title in ['Encapsulation', 'Inheritance']:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = title + ' is a course topic.'
    deck.save(tmp_path / 'deck.pptx')
    records = list(extract(tmp_path / 'deck.pptx', Settings()))
    assert [r['slide'] for r in records] == [1, 2]
    notes = tmp_path / 'notes.md'
    notes.write_text('# Classes\nA class defines objects.\n## Fields\nFields hold state.', encoding='utf-8')
    assert [r['section'] for r in extract(notes, Settings())] == ['Classes', 'Classes / Fields']


def test_corrupt_file_does_not_abort(tmp_path):
    (tmp_path / 'bad.pdf').write_bytes(b'corrupt')
    (tmp_path / 'ok.txt').write_text('Working notes', encoding='utf-8')
    report = ingest(tmp_path, Store(tmp_path / 'test.db'), Settings())
    assert len(report['errors']) == 1
    assert len(report['ingested']) == 1


def test_chunk_location_survives():
    record = dict(page=7, source_type='pdf', text='A grounded sentence. ' * 800)
    chunks = list(chunk_record(record, dict(document_id='d', document_name='a.pdf', original_asset_path='a.pdf')))
    assert len(chunks) > 1
    assert all(c.page == 7 and c.document_name == 'a.pdf' for c in chunks)
    assert len({c.chunk_id for c in chunks}) == len(chunks)
